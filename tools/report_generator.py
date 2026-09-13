"""
report_generator.py

Step 26 of the AI Data Analyst Agent project.

Generates a downloadable report (PDF or Word) summarizing a chat
session: dataset overview, column statistics, any charts created
during the conversation, and the Q&A exchanged with the user.

This is pure Python (fpdf2 for PDF, python-docx for Word) - no LLM
calls happen here. It only formats data that was already computed
elsewhere in the pipeline.
"""

from pathlib import Path
from datetime import datetime

from fpdf import FPDF
from docx import Document
from docx.shared import Inches


REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"


def _format_column_stats(col_name: str, result: dict) -> list:
    """
    Turn one column's analyze_dataset() result into a list of plain
    text lines, formatted based on its type.
    """
    col_type = result.get("type")
    stats = result.get("stats")
    lines = [f"{col_name} ({col_type})"]

    if stats is None:
        lines.append("  No detailed statistics available for this column type.")
        return lines

    if col_type == "numerical":
        lines.append(f"  Mean: {stats['mean']:.2f}   Median: {stats['median']:.2f}")
        lines.append(f"  Min: {stats['min']:.2f}   Max: {stats['max']:.2f}   Std Dev: {stats['std']:.2f}")
        lines.append(f"  Sum: {stats['sum']:.2f}")
    elif col_type == "categorical":
        top_values = ", ".join(f"{k} ({v})" for k, v in stats["top_values"].items())
        lines.append(f"  Most common: {stats['mode']}")
        lines.append(f"  Top values: {top_values}")
    elif col_type == "datetime":
        lines.append(f"  Range: {stats['min_date']} to {stats['max_date']} ({stats['range_days']} days)")

    return lines


def _extract_qa_pairs(messages: list) -> list:
    """
    Pull out (question, answer) pairs from the conversation history,
    skipping system-style messages (dataset loaded confirmation,
    automatic insights) that have no preceding user question.
    """
    pairs = []
    pending_question = None

    for msg in messages:
        if msg["role"] == "user":
            pending_question = msg["content"]
        elif msg["role"] == "assistant" and pending_question is not None:
            pairs.append((pending_question, msg["content"]))
            pending_question = None

    return pairs


def _collect_chart_paths(messages: list) -> list:
    """Pull out every chart image path generated during the conversation."""
    return [msg["image"] for msg in messages if msg.get("image")]


def generate_pdf_report(
    dataset_name: str,
    profile: dict,
    analysis: dict,
    messages: list,
) -> str:
    """
    Build a PDF report and return the path to the saved file.
    """
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "Data Analysis Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"Dataset: {dataset_name}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Dataset Overview", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 7, f"Rows: {profile['rows']}   Columns: {profile['columns']}   Duplicate rows: {profile['duplicate_rows']}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Column Details", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for col_name, result in analysis.items():
        for line in _format_column_stats(col_name, result):
            pdf.multi_cell(0, 6, line)
        pdf.ln(1)
    pdf.ln(4)

    chart_paths = _collect_chart_paths(messages)
    if chart_paths:
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Charts", new_x="LMARGIN", new_y="NEXT")
        for chart_path in chart_paths:
            if Path(chart_path).exists():
                pdf.image(chart_path, w=170)
                pdf.ln(4)

    qa_pairs = _extract_qa_pairs(messages)
    if qa_pairs:
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Questions & Answers", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        for question, answer in qa_pairs:
            pdf.set_font("Helvetica", "B", 10)
            pdf.multi_cell(0, 6, f"Q: {question}")
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, f"A: {answer}")
            pdf.ln(2)

    safe_name = "".join(c if c.isalnum() else "_" for c in dataset_name)
    output_path = REPORTS_DIR / f"report_{safe_name}.pdf"
    pdf.output(str(output_path))

    return str(output_path)


def generate_docx_report(
    dataset_name: str,
    profile: dict,
    analysis: dict,
    messages: list,
) -> str:
    """
    Build a Word (.docx) report and return the path to the saved file.
    """
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()

    doc.add_heading("Data Analysis Report", level=1)
    doc.add_paragraph(f"Dataset: {dataset_name}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_heading("Dataset Overview", level=2)
    doc.add_paragraph(
        f"Rows: {profile['rows']}   Columns: {profile['columns']}   "
        f"Duplicate rows: {profile['duplicate_rows']}"
    )

    doc.add_heading("Column Details", level=2)
    for col_name, result in analysis.items():
        lines = _format_column_stats(col_name, result)
        doc.add_paragraph(lines[0], style="Heading 3")
        for line in lines[1:]:
            doc.add_paragraph(line.strip())

    chart_paths = _collect_chart_paths(messages)
    if chart_paths:
        doc.add_heading("Charts", level=2)
        for chart_path in chart_paths:
            if Path(chart_path).exists():
                doc.add_picture(chart_path, width=Inches(6))

    qa_pairs = _extract_qa_pairs(messages)
    if qa_pairs:
        doc.add_heading("Questions & Answers", level=2)
        for question, answer in qa_pairs:
            q_para = doc.add_paragraph()
            q_para.add_run(f"Q: {question}").bold = True
            doc.add_paragraph(f"A: {answer}")

    safe_name = "".join(c if c.isalnum() else "_" for c in dataset_name)
    output_path = REPORTS_DIR / f"report_{safe_name}.docx"
    doc.save(str(output_path))

    return str(output_path)